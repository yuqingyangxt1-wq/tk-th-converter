"""Dry-run: build a mock EasyBoss xlsx, run convert_source() with TH template, inspect output."""
from __future__ import annotations
import sys, json, os, shutil
from pathlib import Path

import openpyxl
from app.config import default_config
from app.converter import convert_source

# 1) Build mock EasyBoss "导出#SKU" xlsx (PH format — EasyBoss uses uniform format)
mock_src = Path("/tmp/th_dryrun_source.xlsx")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"

headers = [
    "产品名", "品牌", "SKU", "规格1名称", "规格1选项",
    "规格2名称", "规格2选项", "税前价格", "库存",
    "产品图片1", "产品图片2", "产品图片3",
    "包裹重量（KG）", "包裹长度（CM）", "包裹宽度（CM）", "包裹高度（CM）",
    "产品描述", "尺码图",
]
ws.append(headers)

# 1 product, 3 variants (Color=Black, Size=S/M/L)
for size in ["S", "M", "L"]:
    ws.append([
        "COD Unisex T-shirt 现货",         # 产品名
        "",                                # 品牌 (空 → 用 settings)
        f"1640000_abc-Black-{size}",       # SKU
        "颜色",                            # 规格1名称
        "Black",                           # 规格1选项
        "尺码",                            # 规格2名称
        size,                              # 规格2选项
        199,                               # 税前价格 (THB)
        50,                                # 库存
        "https://p16-sg.tiktokcdn.com/obj/example-main.jpg",   # 产品图片1
        "https://p16-sg.tiktokcdn.com/obj/example-2.jpg",
        "https://p16-sg.tiktokcdn.com/obj/example-3.jpg",
        0.2,                               # 包裹重量 KG → 200g
        25, 20, 5,                         # 长宽高 cm
        "Premium cotton t-shirt, soft and breathable. Easy care.",
        "",                                # 尺码图 (空 → 用 settings default)
    ])
wb.save(mock_src)
print(f"mock source saved: {mock_src}  ({mock_src.stat().st_size:,} bytes)")

# 2) Run convert_source
out_dir = Path("/tmp/th_dryrun_out")
if out_dir.exists():
    shutil.rmtree(out_dir)

logs = []
def log(s):
    logs.append(s)
def progress(p, msg):
    logs.append(f"  progress {p:.2f}  {msg}")

settings = default_config()  # fresh config with TH defaults
ps = settings["product_xlsx_settings"]

print()
print("--- DEFAULT_SETTINGS['product_xlsx_settings'] ---")
for k in ["title_prefix", "price_value", "category_value", "brand_value",
           "output_copies", "size_chart_value", "fill_sizes_enabled",
           "split_output_files", "parcel_weight_value"]:
    print(f"  {k} = {ps.get(k)!r}")

template = Path("assets/batch-product-source.xlsx")
result = convert_source(
    source_xlsx=mock_src,
    output_dir=out_dir,
    settings=ps,
    template_src=template,
    progress=progress,
    log=log,
    download_imgs=False,
)

print()
print("--- CONVERT RESULT ---")
print(f"  output_paths: {result.output_paths}")
print(f"  products: {result.product_count}  rows: {result.row_count}")

# 3) Inspect the output xlsx
out_xlsx = result.output_paths[0]
print()
print(f"--- INSPECT {out_xlsx.name} ---")
wb2 = openpyxl.load_workbook(out_xlsx)
ws2 = wb2["Template"]
headers_out = [str(c.value or "") for c in ws2[1]]
print(f"  output column count: {len(headers_out)}")
print(f"  pre_order_time present? {'pre_order_time' in headers_out}")
print(f"  headers: {headers_out}")

print()
print("--- DATA ROWS ---")
for r in range(2, ws2.max_row + 1):
    cells = []
    for c in range(1, ws2.max_column + 1):
        v = ws2.cell(r, c).value
        if v is not None and v != "":
            label = ws2.cell(1, c).value
            cells.append(f"[{label}]={str(v)[:30]}")
    if cells:
        print(f"  R{r}: {' | '.join(cells)}")
    else:
        print(f"  R{r}: (empty)")

# 4) Check key fields
print()
print("--- KEY FIELD CHECK ---")
row1 = ws2[2]
def cell(label):
    idx = headers_out.index(label) + 1
    return ws2.cell(2, idx).value
checks = [
    ("brand", "No brand"),
    ("category", "Men's Tops/T-shirts"),
    ("price", 199),
    ("parcel_weight", 200),  # 0.2 KG → 200g
    ("quantity", 50),  # per-variant stock
    ("seller_sku format", True),
    ("rows total", 6),
]
for label, expected in checks:
    v = cell(label.split()[0])
    if label == "size_chart (non-empty)":
        ok = bool(v)
    elif label == "seller_sku format":
        ok = isinstance(v, str) and "1640000_abc-Black" in v
    elif label == "rows total":
        ok = (ws2.max_row - 1) == expected
        v = ws2.max_row - 1
    elif label == "pre_order_time":
        ok = (v is None)
    else:
        ok = (v == expected)
    flag = "OK" if ok else "FAIL"
    print(f"  [{flag}] {label}: got {v!r}, expected {expected!r}")
